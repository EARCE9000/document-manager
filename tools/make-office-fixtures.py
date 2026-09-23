"""make-office-fixtures.py : Office のテスト用ファイル(test/fixtures/office/)を作る

手書きのXMLではなく実際のアプリが書き出す構造で検証するため、
openpyxl / python-docx / python-pptx を使う。

  pip install openpyxl python-docx python-pptx
  python tools/make-office-fixtures.py test/fixtures/office
"""
import datetime
import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
os.makedirs(OUT, exist_ok=True)

# --- Excel ---
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "売上"
ws.append(["日付", "顧客名", "金額", "担当"])
ws.append([datetime.date(2026, 9, 1), "サンプル商事", 1250000, "山田"])
ws.append([datetime.date(2026, 9, 14), "テスト工業", 830000, "佐藤"])
ws.append([datetime.datetime(2026, 9, 20, 14, 30), "デモ物産", 415000, "鈴木"])
ws["E2"] = "=SUM(C2:C4)"
ws2 = wb.create_sheet("メモ")
ws2["A1"] = "四半期の締めは9月30日"
ws2["A2"] = True
big = wb.create_sheet("大きい表")
for row in range(1, 600):
    big.append([f"行{row}", row * 10, f"備考{row}"])
wb.save(os.path.join(OUT, "sample.xlsx"))

# --- Word ---
from docx import Document

doc = Document()
doc.add_heading("文書管理システム 導入手順書", level=1)
doc.add_paragraph("本書は、社内文書管理システムの導入手順をまとめたものである。")
doc.add_heading("前提条件", level=2)
doc.add_paragraph("サーバーにDockerが導入されていること", style="List Bullet")
doc.add_paragraph("管理者アカウントが払い出されていること", style="List Bullet")
doc.add_heading("手順", level=2)
doc.add_paragraph("設定ファイルを配置し、composeで起動する。")
table = doc.add_table(rows=3, cols=3)
table.cell(0, 0).text = "項番"
table.cell(0, 1).text = "作業"
table.cell(0, 2).text = "担当"
table.cell(1, 0).text = "1"
table.cell(1, 1).text = "環境変数の設定"
table.cell(1, 2).text = "インフラ"
table.cell(2, 0).text = "2"
table.cell(2, 1).text = "疎通確認"
table.cell(2, 2).text = "開発"
doc.save(os.path.join(OUT, "sample.docx"))

# --- PowerPoint ---
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "文書管理システムのご提案"
slide.placeholders[1].text = "2026年9月 情報システム部"

slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "課題"
body = slide.placeholders[1].text_frame
body.text = "資料が個人のPCに散在している"
body.add_paragraph().text = "最新版がどれか分からない"
body.add_paragraph().text = "退職者の資料が引き継がれない"
slide.notes_slide.notes_text_frame.text = "ここで実際の調査結果を紹介する"

slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "効果"
box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
box.text_frame.text = "検索で見つかるようになる"
box.text_frame.add_paragraph().text = "版の取り違えが無くなる"
prs.save(os.path.join(OUT, "sample.pptx"))

for name in ("sample.xlsx", "sample.docx", "sample.pptx"):
    print(name, os.path.getsize(os.path.join(OUT, name)), "bytes")
